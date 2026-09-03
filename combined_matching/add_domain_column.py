"""Add a companyWebsiteDomain column right after dwhId in the spreadsheet.

Looks up each dwhId in Elasticsearch via ElasticCompanyDataSource, and
writes the resolved companyWebsiteDomain into a new column B. Rows whose
domain can't be found are left blank (not removed), so the dwhId sequence
stays intact.
"""

from __future__ import annotations

import sys
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import openpyxl
from dotenv import load_dotenv

load_dotenv()

_ROOT = Path(__file__).resolve().parent.parent
_NM_PATH = str(_ROOT / "name_matching")
if _NM_PATH not in sys.path:
    sys.path.insert(0, _NM_PATH)

from data_source import ElasticCompanyDataSource  # noqa: E402

IN_PATH = Path(__file__).resolve().parent / "Untitled spreadsheet (1).xlsx"
OUT_PATH = IN_PATH  # overwrite in place


def fetch_domain(ds: ElasticCompanyDataSource, dwh_id: str) -> str | None:
    try:
        record = ds.get_by_dwh_id(dwh_id)
    except Exception as exc:
        print(f"  ERROR dwhId={dwh_id}: {exc}", file=sys.stderr)
        return None
    if record is None:
        return None
    domain = record.raw.get("companyWebsiteDomain")
    return domain or None


def main() -> None:
    wb = openpyxl.load_workbook(IN_PATH)
    ws = wb["Sheet1"]

    dwh_ids = [row[0].value for row in ws.iter_rows(min_row=2)]
    print(f"Loaded {len(dwh_ids)} dwhIds")

    ds = ElasticCompanyDataSource()

    results: dict[int, str | None] = {}
    with ThreadPoolExecutor(max_workers=16) as pool:
        futures = {
            pool.submit(fetch_domain, ds, dwh_id): idx
            for idx, dwh_id in enumerate(dwh_ids)
            if dwh_id
        }
        done = 0
        for fut in as_completed(futures):
            idx = futures[fut]
            results[idx] = fut.result()
            done += 1
            if done % 500 == 0:
                print(f"  {done}/{len(futures)} done")

    ws.insert_cols(2)
    ws.cell(row=1, column=2, value="companyWebsiteDomain")

    found = 0
    for idx, dwh_id in enumerate(dwh_ids):
        row = idx + 2
        domain = results.get(idx)
        if domain:
            ws.cell(row=row, column=2, value=domain)
            found += 1

    print(f"Resolved {found}/{len(dwh_ids)} domains")

    wb.save(OUT_PATH)
    print(f"Saved to {OUT_PATH}")


if __name__ == "__main__":
    main()
