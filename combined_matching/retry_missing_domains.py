"""Retry ES lookups for rows left blank in add_domain_column.py's first pass.

The first pass hit transient 'index not found' errors on some dwhIds under
16-way concurrency (same fixed ES_INDEX, so these weren't real misses).
This retries just the blank rows with lower concurrency / more patience.
"""

from __future__ import annotations

import sys
import time
from pathlib import Path

import openpyxl
from dotenv import load_dotenv

load_dotenv()

_ROOT = Path(__file__).resolve().parent.parent
_NM_PATH = str(_ROOT / "name_matching")
if _NM_PATH not in sys.path:
    sys.path.insert(0, _NM_PATH)

from data_source import ElasticCompanyDataSource  # noqa: E402

PATH = Path(__file__).resolve().parent / "Untitled spreadsheet (1).xlsx"


def main() -> None:
    wb = openpyxl.load_workbook(PATH)
    ws = wb["Sheet1"]

    blank_rows = []
    for row in ws.iter_rows(min_row=2):
        dwh_cell, domain_cell = row[0], row[1]
        if dwh_cell.value and not domain_cell.value:
            blank_rows.append((dwh_cell.row, dwh_cell.value))

    print(f"Retrying {len(blank_rows)} blank rows")

    ds = ElasticCompanyDataSource()
    recovered = 0
    still_missing = 0
    for i, (row_num, dwh_id) in enumerate(blank_rows, 1):
        for attempt in range(1, 4):
            try:
                record = ds.get_by_dwh_id(dwh_id)
                domain = record.raw.get("companyWebsiteDomain") if record else None
                break
            except Exception as exc:
                domain = None
                if attempt < 3:
                    time.sleep(3)
                else:
                    print(f"  still failing dwhId={dwh_id}: {exc}", file=sys.stderr)

        if domain:
            ws.cell(row=row_num, column=2, value=domain)
            recovered += 1
        else:
            still_missing += 1

        if i % 100 == 0:
            print(f"  {i}/{len(blank_rows)} processed, recovered {recovered}")

    print(f"Recovered {recovered}, still missing {still_missing}")
    wb.save(PATH)
    print("Saved.")


if __name__ == "__main__":
    main()
